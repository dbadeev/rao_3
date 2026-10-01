#!/usr/bin/env python3
"""slide_index.py - slide number <-> JSON index table for a PPTX (index is 0-based; slide_number = index + 1).
Usage: python slide_index.py template.pptx -o slide_index.csv
Note: fingerprint covers layout name and top-level shape ids/names/types/text, NOT geometry.
"""
import argparse, csv, hashlib
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def fingerprint(slide):
    parts = [slide.slide_layout.name]
    for s in slide.shapes:
        text = s.text_frame.text.strip()[:40] if s.has_text_frame else ""
        parts.append(f"{s.shape_id}|{s.name}|{str(s.shape_type)}|{text}")
    return hashlib.sha1("\n".join(parts).encode("utf-8")).hexdigest()[:12]

def first_text(slide):
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip():
            return s.text_frame.text.strip().replace("\n", " / ")[:70]
    return ""

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx"); ap.add_argument("-o", "--output", default="slide_index.csv")
    a = ap.parse_args()
    prs = Presentation(a.pptx); rows = []
    for i, sl in enumerate(prs.slides):
        sh = list(sl.shapes)
        rows.append({"slide_number": i + 1, "json_index": i, "layout": sl.slide_layout.name, "n_shapes": len(sh),
            "n_text": sum(s.has_text_frame and bool(s.text_frame.text.strip()) for s in sh),
            "n_tables": sum(s.has_table for s in sh), "n_charts": sum(s.has_chart for s in sh),
            "n_pictures": sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in sh),
            "first_text": first_text(sl), "fingerprint": fingerprint(sl)})
    with open(a.output, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    print(f"Slides: {len(rows)}; CSV: {a.output}")
if __name__ == "__main__": main()
