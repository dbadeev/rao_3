"""Automatic per-slide template catalog: deterministic inventory plus optional text/VLM adjudication.
The input PPTX is never modified. Shape references are stable within the input file (slide index + shape ID).
"""
import argparse, base64, io, json, os, re, subprocess, tempfile, zipfile
from collections import Counter
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TYPES=("title","section","slide","table","last")
ROLES=("TITLE_FIELD","SUBTITLE_FIELD","TEXT_FIELD","TABLE_FIELD","IMAGE_FIELD")
EMU=914400

def text_of(s): return s.text_frame.text.strip() if s.has_text_frame else ""
def info(s):
 return {"id":s.shape_id,"name":s.name,"kind":str(s.shape_type).split(" (")[0],
         "placeholder_kind":str(s.placeholder_format.type).split(" (")[0] if s.is_placeholder else None,
         "placeholder_idx":s.placeholder_format.idx if s.is_placeholder else None,
         "text":text_of(s)[:240],"x":round(s.left/EMU,2),"y":round(s.top/EMU,2),
         "w":round(s.width/EMU,2),"h":round(s.height/EMU,2)}

def inventory(path):
 prs=Presentation(path); slides=[]
 for i,sl in enumerate(prs.slides):
  shapes=[info(s) for s in sl.shapes]
  slides.append({"index":i,"layout_name":sl.slide_layout.name,"shapes":shapes,
                 "summary":{"n_tables":sum(s.has_table for s in sl.shapes),"n_charts":sum(s.has_chart for s in sl.shapes),
                            "n_pictures":sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in sl.shapes)}})
 return {"source":Path(path).name,"total_slides":len(slides),"slide_size_inches":[round(prs.slide_width/EMU,3),round(prs.slide_height/EMU,3)],"slides":slides}

def heuristic(item, rank, last):
 sh=item['shapes']; named=item['layout_name'].lower(); summary=item['summary']
 blank=not any(x['text'] or x['kind'] in ('PICTURE','TABLE','CHART') for x in sh)
 if rank==0: typ='title'; score=.9
 elif rank==last: typ='last'; score=.7
 elif blank and ('раздел' in named or 'section' in named): typ='section';score=.7
 elif blank: typ='skip';score=.9
 elif summary['n_tables']:typ='table';score=.96
 elif ('раздел' in named or 'section' in named) and not summary['n_charts'] and not summary['n_pictures']:typ='section';score=.82
 elif len([x for x in sh if x['text']])<=2 and len(sh)<=5 and not summary['n_charts'] and not summary['n_pictures']:typ='section';score=.5
 else:typ='slide';score=.72
 fields=[]; used=set()
 def add(role,x,confidence,reason):
  if x is None or x['id'] in used:return
  fields.append({'role':role,'shape_id':x['id'],'confidence':confidence,'reason':reason});used.add(x['id'])
 def candidates(role):
  if role=='TABLE_FIELD':return [x for x in sh if x['kind']=='TABLE' or x['placeholder_kind']=='TABLE']
  if role=='IMAGE_FIELD':return [x for x in sh if x['kind']=='PICTURE' or x['placeholder_kind']=='PICTURE']
  return [x for x in sh if x['text'] or x['placeholder_kind'] in ('TITLE','CENTER_TITLE','BODY','SUBTITLE')]
 wanted=['TITLE_FIELD']+(['SUBTITLE_FIELD'] if typ in ('title','section') else [])+(['TEXT_FIELD','IMAGE_FIELD'] if typ in ('slide','table') else [])+(['TABLE_FIELD'] if typ=='table' else [])
 ph_map={'TITLE_FIELD':('TITLE','CENTER_TITLE'),'SUBTITLE_FIELD':('SUBTITLE',),'TEXT_FIELD':('BODY',),'TABLE_FIELD':('TABLE',),'IMAGE_FIELD':('PICTURE',)}
 for role in wanted:
  arr=[x for x in candidates(role) if x['id'] not in used]
  exact=next((x for x in arr if x['name'].upper()==role),None)
  ph=next((x for x in arr if x['placeholder_kind'] in ph_map[role]),None)
  if exact:add(role,exact,1.,'exact_name');continue
  if ph:add(role,ph,.94,'placeholder_type');continue
  if role in ('TABLE_FIELD','IMAGE_FIELD'):
   if len(arr)==1:add(role,arr[0],.85,'unique_native_shape')
   continue
  if role=='TITLE_FIELD':
   arr.sort(key=lambda x:(0 if re.search('заголов|title',x['name'],re.I) else 1,x['y']))
  elif role=='SUBTITLE_FIELD':arr.sort(key=lambda x:(0 if re.search('подзаголов|subtitle',x['name'],re.I) else 1,x['y']))
  else:arr=[x for x in arr if x['y']>.7];arr.sort(key=lambda x:-len(x['text']))
  if arr:add(role,arr[0],.5,'ambiguous_geometry')
 warnings=[]
 # if typ=='section' and not any(f['role']=='SUBTITLE_FIELD' for f in fields):warnings.append('section_subtitle_unresolved')
 if typ in ('slide','table') and not any(f['role']=='TEXT_FIELD' for f in fields):warnings.append('text_field_unresolved')
 if any(f['confidence']<.7 for f in fields) or score<.75:warnings.append('low_confidence')
 return {'type':typ,'confidence':score,'fields':fields,'warnings':warnings}

def parse_json(raw):
 raw=raw.strip();raw=re.sub(r'^```(?:json)?\s*|\s*```$','',raw,flags=re.I)
 try:return json.loads(raw)
 except json.JSONDecodeError:
  decoder=json.JSONDecoder();start=raw.find('{')
  if start>=0:return decoder.raw_decode(raw[start:])[0]
  raise

def review(item, model, base_url, api_key, image=None):
 from openai import OpenAI
 client=OpenAI(base_url=base_url,api_key=api_key,timeout=300,max_retries=1)
 visible=[s for s in item['shapes'] if s['text'] or s['kind'] in ('PICTURE','TABLE','CHART')]
 short=[{k:s[k] for k in ('id','name','kind','placeholder_kind','text','x','y','w','h')} for s in visible[:70]]
 prompt=("Ответь только JSON: {\"type\":\"title|section|slide|table|last\", \"fields\":[{\"role\":\"TITLE_FIELD|SUBTITLE_FIELD|TEXT_FIELD|TABLE_FIELD|IMAGE_FIELD\",\"shape_id\":123}],\"reason\":\"...\"}. "
         "Оцени русский PPTX слайд как многоразовый образец. A section slide may have no subtitle. Assign SUBTITLE_FIELD only if a separate"
         " editable subtitle shape actually exists. Never invent shape IDs. "
         "Выбирай ТОЛЬКО id существующих редактируемых фигур; групповые и неописанные фигуры не выдумывай. "
         "Не принимай логотип/фон за IMAGE_FIELD. Если картинка слайд-скриншот, отметь ограничения в reason. "
         +json.dumps({'index':item['index'],'layout':item['layout_name'], 'shapes':short},ensure_ascii=False))
 content=[{'type':'text','text':prompt}]
 if image:
  content.append({'type':'image_url','image_url':{'url':'data:image/jpeg;base64,'+base64.b64encode(image).decode()}})
 answer=client.chat.completions.create(model=model,messages=[{'role':'user','content':content if image else prompt}],temperature=0,max_tokens=550)
 return parse_json(answer.choices[0].message.content)

def validate(reviewed,item):
 ids={s['id']:s for s in item['shapes']}; typ=reviewed.get('type')
 if typ not in TYPES:raise ValueError('invalid slide type')
 fields=[]; seen=set(); used=set()
 for f in reviewed.get('fields',[]):
  role=f.get('role');sid=f.get('shape_id');sid=int(sid) if str(sid).isdigit() else sid
  if role not in ROLES or role in seen or sid not in ids or sid in used:raise ValueError('invalid/duplicate field reference')
  shape=ids[sid]
  if role=='TABLE_FIELD' and shape['kind']!='TABLE' and shape['placeholder_kind']!='TABLE':raise ValueError('table field is not native table')
  if role=='IMAGE_FIELD' and shape['kind']!='PICTURE' and shape['placeholder_kind']!='PICTURE':raise ValueError('image field is not picture')
  if role in ('TITLE_FIELD','SUBTITLE_FIELD','TEXT_FIELD') and shape['kind'] not in ('TEXT_BOX','AUTO_SHAPE','PLACEHOLDER'):raise ValueError('not editable text')
  fields.append({'role':role,'shape_id':sid,'confidence':None,'reason':'model_validated_reference'});seen.add(role);used.add(sid)
 warnings=[]
 # if typ=='section' and 'SUBTITLE_FIELD' not in seen:warnings.append('section_subtitle_unresolved')
 if typ=='table' and 'TABLE_FIELD' not in seen:warnings.append('table_field_unresolved')
 return {'type':typ,'confidence':None,'fields':fields,'warnings':warnings,'reason':str(reviewed.get('reason',''))[:300]}


def review_status(item):
 sel = item['selected']
 if sel['type'] == 'skip':
  return 'not_applicable'
 if 'model_error' in item or 'low_confidence' in sel.get('warnings', []):
  return 'needs_human_review'
 shapes = item['shapes']
 n_text = sum(1 for s in shapes if s['text'])
 complex_slide = n_text > 6 or any(s['kind'] in ('CHART', 'GROUP') for s in shapes)
 return 'needs_human_review' if complex_slide else 'unreviewed'

def render(path):
 from PIL import Image
 from pdf2image import convert_from_path
 with tempfile.TemporaryDirectory() as tmp:
  profile=Path(tmp)/'profile'; out=Path(tmp)/'pdf';out.mkdir()
  p=subprocess.run(['soffice','-env:UserInstallation=file://'+str(profile),'--headless','--convert-to','pdf','--outdir',str(out),str(Path(path).resolve())],capture_output=True,text=True,timeout=180)
  pdf=out/(Path(path).stem+'.pdf')
  if p.returncode or not pdf.exists():raise RuntimeError(p.stderr or p.stdout)
  pages=convert_from_path(str(pdf),dpi=110)
  from PIL import ImageOps
  images=[]
  for page in pages:
   rgb=page.convert('RGB');rgb.thumbnail((1600,1000))
   buffer=io.BytesIO();rgb.save(buffer,format='JPEG',quality=82);images.append(buffer.getvalue())
  return images


def run(path,mode,base_url=None,model=None,api_key=None,max_calls=0):
 deck=inventory(path);slides=deck['slides']; active=[s['index'] for s in slides if any(x['text'] or x['kind'] in ('PICTURE','TABLE','CHART') for x in s['shapes'])]
 active=sorted(set(active+[0,len(slides)-1])) if slides else []
 for item in slides:
  rank=active.index(item['index']) if item['index'] in active else -1
  item['heuristic']=heuristic(item,rank,len(active)-1)
 images=render(path) if mode=='vlm' else []
 count=0
 for item in slides:
  if mode=='heuristic' or item['heuristic']['type']=='skip':continue
  if max_calls and count>=max_calls:break
  count+=1
  try:
   img=images[item['index']] if mode=='vlm' else None
   proposal=review(item,model,base_url,api_key,image=img)
   item['model_proposal_raw']=proposal
   item['model_validated']=validate(proposal,item)
   item['selected']=item['model_validated']
   item['selected_source'] = 'model'
  except Exception as e:item['model_error']=str(e)
 for item in slides:
  item.setdefault('selected',item['heuristic'])
  item.setdefault('selected_source', 'heuristic')
  item['review_status'] = review_status(item)
  if mode=='vlm' and item['index']>=len(images):item.setdefault('model_error','render_missing')
 deck.update({'mode':mode,'model':model if mode!='heuristic' else None,'reviewed_count':count,
              'notes':['No PPTX file is modified. Candidate catalog is not a rewritten Slide Master.',
                       'Selected fields refer to source slide shape IDs; validate before injection.']})
 return deck

def main():
 p=argparse.ArgumentParser();p.add_argument('pptx');p.add_argument('--mode',choices=['heuristic','text','vlm'],default='heuristic');p.add_argument('--output',required=True)
 p.add_argument('--base-url',default=os.getenv('LLM_BASE_URL'));p.add_argument('--model',default=os.getenv('LLM_MODEL'));p.add_argument('--max-calls',type=int,default=0)
 a=p.parse_args()
 if a.mode!='heuristic' and (not a.base_url or not a.model):p.error('text/vlm requires --base-url and --model')
 result=run(a.pptx,a.mode,a.base_url,a.model,os.getenv('LLM_API_KEY','EMPTY'),a.max_calls)
 Path(a.output).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'slides':result['total_slides'],'types':dict(Counter(x['selected']['type'] for x in result['slides'])),'output':a.output,'model_errors':sum('model_error' in x for x in result['slides'])},ensure_ascii=False))
if __name__=='__main__':main()
