from pptx import Presentation
from pptx.oxml.ns import qn
FIELD_NAMES = ("TITLE_FIELD", "SUBTITLE_FIELD", "TEXT_FIELD", "TABLE_FIELD", "IMAGE_FIELD")
DEFAULT_TABLE_STYLE_ID = "{5DA37D80-6434-44D0-A028-1B22A696006F}"

def inspect_template(path):
    prs = Presentation(path)
    report = {"slide_width_in": prs.slide_width / 914400, "slide_height_in": prs.slide_height / 914400, "layouts": []}
    for layout in prs.slide_layouts:
        fields = {}
        for ph in layout.placeholders:
            if ph.name in FIELD_NAMES:
                fields[ph.name] = ph.placeholder_format.idx
        report["layouts"].append({"name": layout.name, "fields": fields})
    return report

def build_layout_field_map(layout):
    mapping = {}
    for ph in layout.placeholders:
        if ph.name in FIELD_NAMES:
            mapping[ph.placeholder_format.idx] = ph.name
    return mapping

def get_field_by_role(slide, field_map, role):
    for ph in slide.placeholders:
        if field_map.get(ph.placeholder_format.idx) == role:
            return ph
    return None

def apply_table_style(table, style_id=DEFAULT_TABLE_STYLE_ID, first_row=True, band_row=True):
    tbl_elm = table._tbl
    tblPr = tbl_elm.find(qn('a:tblPr'))
    if tblPr is None:
        tblPr = tbl_elm.makeelement(qn('a:tblPr'), {})
        tbl_elm.insert(0, tblPr)
    tblPr.set('firstRow', '1' if first_row else '0')
    tblPr.set('bandRow', '1' if band_row else '0')
    for child in list(tblPr):
        tblPr.remove(child)
    style_elm = tblPr.makeelement(qn('a:tableStyleId'), {})
    style_elm.text = style_id
    tblPr.append(style_elm)

def choose_layout_name(item):
    t = item["slide_type"]
    if t == "title": return "title1"
    if t == "closing": return "last1"
    if t == "section": return "section1"
    if t == "kpi_table": return "table_text1" if item.get("kpis") else "table_only1"
    if t == "timeline": return "table_only1"
    if t == "bullets" and item.get("image_path"): return "slide2"
    return "slide1"

def set_title(slide, field_map, text):
    ph = get_field_by_role(slide, field_map, "TITLE_FIELD")
    if ph and text: ph.text_frame.text = text

def set_bullets(slide, field_map, bullets):
    ph = get_field_by_role(slide, field_map, "TEXT_FIELD")
    if not ph or not bullets: return
    tf = ph.text_frame
    tf.text = bullets[0]
    for b in bullets[1:]:
        p = tf.add_paragraph(); p.text = b

def set_table(slide, field_map, table_data, style_id=DEFAULT_TABLE_STYLE_ID):
    ph = get_field_by_role(slide, field_map, "TABLE_FIELD")
    if not ph or not table_data: return
    rows = len(table_data["rows"]) + 1
    cols = len(table_data["headers"])
    graphic_frame = ph.insert_table(rows, cols)
    table = graphic_frame.table
    apply_table_style(table, style_id=style_id)
    for c, h in enumerate(table_data["headers"]):
        table.cell(0, c).text = h
    for r, row in enumerate(table_data["rows"], start=1):
        for c, val in enumerate(row):
            table.cell(r, c).text = str(val)

def set_image(slide, field_map, image_path):
    ph = get_field_by_role(slide, field_map, "IMAGE_FIELD")
    if ph and image_path: ph.insert_picture(image_path)

def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].get(qn('r:id'))
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])

def build_presentation(manifest, template_path, out_path, drop_sample_slides=True, table_style_id=DEFAULT_TABLE_STYLE_ID):
    prs = Presentation(template_path)
    n_existing = len(prs.slides)
    layouts = {l.name: l for l in prs.slide_layouts}
    field_maps = {name: build_layout_field_map(l) for name, l in layouts.items()}
    for item in manifest["slides"]:
        layout_name = choose_layout_name(item)
        if layout_name not in layouts:
            raise ValueError(f"В шаблоне не найден layout '{layout_name}'. Доступные: {list(layouts)}")
        layout = layouts[layout_name]
        slide = prs.slides.add_slide(layout)
        fmap = field_maps[layout_name]
        if item["slide_type"] in ("title", "section", "closing"):
            set_title(slide, fmap, item.get("heading") or item.get("closing_text") or manifest.get("deck_title", ""))
        elif item["slide_type"] == "kpi_table":
            set_title(slide, fmap, item["heading"])
            kpis = item.get("kpis") or []
            set_bullets(slide, fmap, [f"{k['label']}: {k['value']}" + (f" ({k['delta']})" if k.get('delta') else "") for k in kpis])
            set_table(slide, fmap, item.get("table"), style_id=table_style_id)
        elif item["slide_type"] == "timeline":
            set_title(slide, fmap, item["heading"])
            set_table(slide, fmap, {"headers": ["Дата", "Событие"], "rows": [[e["date"], e["label"]] for e in item["timeline"]]}, style_id=table_style_id)
        elif item["slide_type"] == "bullets":
            set_title(slide, fmap, item["heading"])
            set_bullets(slide, fmap, item.get("bullets", []))
            set_image(slide, fmap, item.get("image_path"))
    if drop_sample_slides:
        for idx in reversed(range(n_existing)):
            delete_slide(prs, idx)
    prs.save(out_path)
    return prs
