#!/usr/bin/env python3
"""check_gold.py - validate a labels JSON against the PPTX (shape ids may point into groups).
Usage: python check_gold.py template.pptx labels.json   (slide_index.py and shapes_util.py must be in the same folder)
Author's rules checked here:
  - a BLANK slide (case["blank"] == true) has no labelled fields at all (it is the "zero" element of a deck);
  - an EMPTY text placeholder on a non-blank slide IS labelled (the field is where the title goes); these are listed as INFO;
  - SUBTITLE_FIELD outside title/section is a WARNING (README-2.md defines it only for title*/section*).
"""
import json, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from slide_index import fingerprint
from shapes_util import index_shapes

TYPES = {"title", "section", "slide", "table", "last"}
ROLES = {"TITLE_FIELD", "SUBTITLE_FIELD", "TEXT_FIELD", "TABLE_FIELD", "IMAGE_FIELD"}
README_ALLOWED = {"title": {"TITLE_FIELD", "SUBTITLE_FIELD"}, "section": {"TITLE_FIELD", "SUBTITLE_FIELD"},
    "slide": {"TITLE_FIELD", "TEXT_FIELD", "IMAGE_FIELD"}, "table": {"TITLE_FIELD", "TEXT_FIELD", "IMAGE_FIELD", "TABLE_FIELD"}, "last": {"TITLE_FIELD"}}
USABILITY = {"full", "partial", "static", "reference_only"}

def main(pptx, gold_path):
    prs = Presentation(pptx); gold = json.load(open(gold_path, encoding="utf-8")); errors, warnings, info = [], [], []
    for c in gold["cases"]:
        cid = c.get("case_id", "?"); idx = c.get("json_index")
        if not isinstance(idx, int) or not 0 <= idx < len(prs.slides):
            errors.append(f"{cid}: json_index outside 0..{len(prs.slides)-1}"); continue
        slide = prs.slides[idx]; by_id = {k: v[0] for k, v in index_shapes(slide).items()}
        if c.get("fingerprint") != fingerprint(slide): errors.append(f"{cid}: fingerprint mismatch")
        if c.get("slide_number") != idx + 1: errors.append(f"{cid}: slide_number must equal json_index+1")
        if c.get("type") not in TYPES: errors.append(f"{cid}: invalid type {c.get('type')!r}")
        if c.get("usability") not in USABILITY: errors.append(f"{cid}: invalid usability")
        if c.get("selection") not in ("included", "excluded"): errors.append(f"{cid}: selection must be included/excluded")
        used = {}
        if c.get("blank") and any(c["fields"].values()): errors.append(f"{cid}: blank slide must have no labelled fields")
        for role, ids in c.get("fields", {}).items():
            if role not in ROLES: errors.append(f"{cid}: unknown role {role}"); continue
            if ids and c.get("type") in README_ALLOWED and role not in README_ALLOWED[c["type"]]:
                warnings.append(f"{cid} (slide {idx+1}): {role} on type {c['type']} is an extension of README-2.md")
            for sid in ids or []:
                s = by_id.get(sid)
                if s is None: errors.append(f"{cid}: {role} shape_id {sid} absent on slide {idx+1}"); continue
                if sid in used: errors.append(f"{cid}: shape {sid} assigned to both {used[sid]} and {role}")
                used[sid] = role
                if role == "TABLE_FIELD" and not s.has_table: errors.append(f"{cid}: TABLE_FIELD {sid} is not a table")
                if role == "IMAGE_FIELD" and not (s.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE) or s.has_chart):
                    errors.append(f"{cid}: IMAGE_FIELD {sid} has unsuitable shape type")
                if role in {"TITLE_FIELD", "SUBTITLE_FIELD", "TEXT_FIELD"}:
                    if not s.has_text_frame: errors.append(f"{cid}: {role} {sid} has no text frame")
                    elif not s.text_frame.text.strip(): info.append(f"{cid} (slide {idx+1}): {role} id{sid} is an empty placeholder, labelled on purpose")
        for pr in c.get("paragraph_roles", []):
            s = by_id.get(pr.get("shape_id"))
            if s is None or not s.has_text_frame or pr["shape_id"] not in c["fields"].get("TEXT_FIELD", []): errors.append(f"{cid}: bad paragraph_roles {pr}")
        for cp in c.get("compound_images", []):
            if cp["shape_id"] not in c["fields"].get("IMAGE_FIELD", []) or any(p not in by_id or p in used for p in cp["parts"]): errors.append(f"{cid}: bad compound_images {cp}")
    inc = sum(c["selection"] == "included" for c in gold["cases"])
    print(f"Cases: {len(gold['cases'])} (included {inc}); errors: {len(errors)}; warnings: {len(warnings)}; info: {len(info)}")
    for x in errors: print("ERROR  ", x)
    for x in warnings: print("WARNING", x)
    for x in info: print("INFO   ", x)
    raise SystemExit(1 if errors else 0)
if __name__ == "__main__": main(sys.argv[1], sys.argv[2])
