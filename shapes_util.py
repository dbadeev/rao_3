"""shapes_util.py - index every shape of a slide, including children of groups, with ABSOLUTE geometry (EMU)."""
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

def _walk(shapes, tx, out, top_id):
    for s in shapes:
        x = tx[0] + (s.left - tx[2]) * tx[4]; y = tx[1] + (s.top - tx[3]) * tx[5]
        box = (x, y, s.width * tx[4], s.height * tx[5])
        out.setdefault(s.shape_id, (s, box, top_id if top_id is not None else s.shape_id))
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            xf = s._element.grpSpPr.find(qn("a:xfrm"))
            off, ext = xf.find(qn("a:off")), xf.find(qn("a:ext"))
            cho, che = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
            sx = int(ext.get("cx")) / max(int(che.get("cx")), 1) * tx[4]; sy = int(ext.get("cy")) / max(int(che.get("cy")), 1) * tx[5]
            gx = tx[0] + (int(off.get("x")) - tx[2]) * tx[4]; gy = tx[1] + (int(off.get("y")) - tx[3]) * tx[5]
            _walk(s.shapes, (gx, gy, int(cho.get("x")), int(cho.get("y")), sx, sy), out, top_id if top_id is not None else s.shape_id)

def index_shapes(slide):
    """-> {shape_id: (shape, (x, y, w, h) in EMU on the slide, id_of_top_level_ancestor)}"""
    out = {}
    _walk(slide.shapes, (0, 0, 0, 0, 1.0, 1.0), out, None)
    return out
