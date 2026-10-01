#!/usr/bin/env python3
"""build_labels.py - labels for Shablon-prezentatsii_RUS_16x9.pptx (28 slides), version 8 (author's decisions of 1 Oct).
Usage: python build_labels.py template.pptx labels.json      (also writes <name>_selected.json with the included slides only)
Decisions: slide 3 is a BLANK slide ("zero" of a deck): kept in the set, nothing labelled; an empty title placeholder on any other slide IS labelled;
slide 25 is included again; 13 slides are excluded (diagrams, numbers, complex composition).
Rules:
  - IMAGE_FIELD: top-level PICTURE, CHART, icon groups (1.3x1.3 in of AUTO_SHAPE/PICTURE), children of groups made only of PICTURE,
    plus IMAGE_EXTRA (loose icons labelled by the OUTER circle; the glyph inside is in compound_images).
  - SUBTITLE_FIELD: explicit SUBTITLES (an extension of README-2.md for slide/table).
  - TEXT_FIELD: shapes with 'Образец текста', reading order; a bold first heading paragraph inside is recorded in paragraph_roles.
"""
import json, re, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from slide_index import fingerprint
from shapes_util import index_shapes

EXCLUDED = {12, 13, 14, 15, 16, 18, 19, 20, 21, 23, 24, 26, 27}          # 1-based slide numbers (author's selection)
BLANK = {2}                                                                # json_index of blank slides (nothing is labelled)
SUBTITLES = {8: [38], 9: [17, 50, 28, 34], 10: [21, 22, 23, 30], 16: [46], 17: [33], 18: [27], 21: [17, 81925, 3, 21], 24: [7, 8]}
IMAGE_EXTRA = {9: [31], 10: [20]}
IMAGE_EXCLUDE = {10: [32]}
COMPOUND = {9: [{"shape_id": 31, "parts": [41], "note": "Значок «земной шар»: круг id31 + контур id41"}],
            10: [{"shape_id": 20, "parts": [32], "note": "Значок: круг id20 + картинка id32 внутри"}]}
TYPE_SPEC = {0: ("title", [], "full"), 1: ("section", [], "full"), 2: ("slide", [], "static"), 3: ("slide", [], "full"), 4: ("slide", [], "full"),
 5: ("slide", [], "reference_only"), 6: ("section", ["title"], "full"), 7: ("table", ["slide"], "partial"), 8: ("table", ["slide"], "partial"),
 23: ("slide", [], "partial"), 25: ("slide", [], "reference_only"), 26: ("table", ["slide"], "partial"), 27: ("last", [], "static")}
UNCERTAIN = {26}
NOTES = {
 0: {"SUBTITLE_FIELD": "Блок id3 содержит докладчика, должность и дату одним текстом"},
 1: {"SUBTITLE_FIELD": "Отдельной фигуры подзаголовка нет; поле не обязательно", "TITLE_FIELD": "Заполнитель пуст, но выделен цветом макета и является местом заголовка: размечен"},
 2: {"TITLE_FIELD": "Пустой слайд — «ноль» набора: ничего не размечается и не заполняется; вставляется, если текст пользователя требует пустого слайда"},
 3: {"TITLE_FIELD": "Пустой заполнитель размечен: ниже содержательный объект (картинка)", "TEXT_FIELD": "Текстовой области нет"},
 4: {"TITLE_FIELD": "Пустой заполнитель размечен: ниже содержательный объект (картинка)", "TEXT_FIELD": "Текстовой области нет"},
 5: {"TEXT_FIELD": "Справка по шрифтам и оформлению цифр, не содержательный образец"},
 6: {"SUBTITLE_FIELD": "Блок id4 (вопрос, название вопроса, докладчик) принят как один подзаголовок"},
 7: {"TEXT_FIELD": "id803 — подпись диаграммы; плашка «Важная информация» не размечена: у table текст только с буллетами",
     "IMAGE_FIELD": "Рамка диаграммы id25 в исходнике заходит на таблицу id28 (по высоте на ~0,3 дюйма)"},
 8: {"TITLE_FIELD": "Заголовок — обычный текстовый блок, не placeholder", "SUBTITLE_FIELD": "Расширение README: плашка id38"},
 9: {"IMAGE_FIELD": "Значок «земной шар» размечен внешним кругом id31; контур id41 входит в значок (compound_images)"},
 10: {"IMAGE_FIELD": "Значок id20: размечен внешний круг; картинка id32 внутри входит в значок (compound_images)"},
 21: {"SUBTITLE_FIELD": "id3 «ПОДЗАГОЛОВОК:» лежит поверх верхней части сохранённой рамки блока id16 (общий верх 2,77 дюйма)"},
 23: {"TEXT_FIELD": "Слайд состоит из диаграмм, подписей и значений; текста с буллитами нет"},
 24: {"SUBTITLE_FIELD": "Расширение README: «НАЗВАНИЕ БЛОКА» id7 и id8"},
 25: {"TEXT_FIELD": "175 фигур: диаграммы-водопады, подписи и значения; текстовой области нет"},
 26: {"TEXT_FIELD": "Текстовой области нет; кроме таблицы — 3 диаграммы и десятки подписей"},
 27: {"TITLE_FIELD": "Надпись «СПАСИБО ЗА ВНИМАНИЕ!» закреплена на макете, фигур на слайде нет"},
}
HEAD_RE = re.compile(r"(заголов|подзаголов|название\s+(блока|графика|темы)|^блок\s*\d|направления)", re.I)
def kind(s): return str(s.shape_type).split(" (")[0]

def main(src, dst):
    prs = Presentation(src); H = prs.slide_height; cases = []
    for i, sl in enumerate(prs.slides):
        idx = index_shapes(sl); top = list(sl.shapes)
        ph = [s.shape_id for s in top if s.is_placeholder and "TITLE" in str(s.placeholder_format.type)]
        title = {0: [2], 6: [6], 8: [21], 27: []}.get(i, ph[:1])
        if i in BLANK: title = []
        body = [s for s in top if not s.is_placeholder and s.has_text_frame and "Образец текста" in s.text_frame.text]
        body.sort(key=lambda s: (round(s.top / H, 1), s.left))
        images = []
        for s in top:
            kids = list(s.shapes) if s.shape_type == MSO_SHAPE_TYPE.GROUP else []
            icon = kids and abs(s.width/914400-1.3) < .06 and abs(s.height/914400-1.3) < .06 and {kind(c) for c in kids} <= {"AUTO_SHAPE", "PICTURE"}
            if kids and {kind(c) for c in kids} == {"PICTURE"}: images += [c.shape_id for c in kids]
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE or s.has_chart or icon: images.append(s.shape_id)
        images = [k for k in images if k not in IMAGE_EXCLUDE.get(i, [])] + IMAGE_EXTRA.get(i, [])
        images = sorted(set(images), key=lambda k: (round(idx[k][1][1] / H, 1), idx[k][1][0]))
        typ, alts, use = TYPE_SPEC.get(i, ("slide", [], "partial"))
        fields = {"TITLE_FIELD": title}
        if typ in ("title", "section"): fields["SUBTITLE_FIELD"] = {0: [3], 6: [4]}.get(i, [])
        elif typ != "last":
            if i in SUBTITLES:
                fields["SUBTITLE_FIELD"] = sorted(SUBTITLES[i], key=lambda k: (round(idx[k][1][1] / H, 1), idx[k][1][0])) if i in (21, 24) else SUBTITLES[i]
            fields["TEXT_FIELD"] = [] if i in (2, 3, 4, 5, 7, 23, 25, 26) else ([37] if i == 8 else [s.shape_id for s in body])
        if typ == "table": fields["TABLE_FIELD"] = [s.shape_id for s in top if s.has_table]
        if typ in ("slide", "table"): fields["IMAGE_FIELD"] = [] if i in BLANK else images
        notes = dict(NOTES.get(i, {}))
        if fields.get("TEXT_FIELD") and "TEXT_FIELD" not in notes:
            notes["TEXT_FIELD"] = "Блоки с образцом текста слева направо, сверху вниз; остальные образцы на слайде не размечены"
        para = []
        for sid in fields.get("TEXT_FIELD", []):
            s = idx[sid][0]; ps = s.text_frame.paragraphs
            if len(ps) >= 2 and "Образец текста" not in ps[0].text and ps[0].runs and ps[0].runs[0].font.bold:
                para.append({"shape_id": sid, "heading_paragraphs": 1, "heading_role": "SUBTITLE_FIELD", "body_role": "TEXT_FIELD"})
        labelled = {k for ids in fields.values() for k in (ids or [])}
        open_c = [s.shape_id for s in top if not s.is_placeholder and s.has_text_frame and s.shape_id not in labelled
                  and s.text_frame.text.strip() and len(s.text_frame.text.strip().split("\n")[0]) <= 60
                  and HEAD_RE.search(s.text_frame.text.strip().split("\n")[0]) and "образец текста" not in s.text_frame.text.lower()]
        sel = "excluded" if i + 1 in EXCLUDED else "included"
        cases.append({"case_id": f"d{i+1:02d}", "slide_number": i + 1, "json_index": i, "fingerprint": fingerprint(sl), "layout": sl.slide_layout.name,
            "selection": sel, "selection_reason": "диаграммы, цифры или сложный состав (решение автора)" if sel == "excluded" else "",
            "blank": i in BLANK, "type": typ, "type_alternatives": alts, "usability": use, "fields": fields, "paragraph_roles": para,
            "compound_images": COMPOUND.get(i, []), "open_candidates": open_c, "field_notes": notes, "uncertain": i in UNCERTAIN,
            "review": {"status": "v8_author_decisions_1_oct", "author": "assistant + author remarks"}})
    json.dump({"source": "Shablon-prezentatsii_RUS_16x9.pptx (28 slides)", "cases": cases}, open(dst, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    inc = [c["slide_number"] for c in cases if c["selection"] == "included"]
    json.dump({"source": "Shablon-prezentatsii_RUS_16x9.pptx (28 slides), selected slides only", "cases": [c for c in cases if c["selection"] == "included"]},
              open(dst.replace(".json", "_selected.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print("cases:", len(cases), "included:", len(inc), inc)
if __name__ == "__main__": main(*sys.argv[1:3])
